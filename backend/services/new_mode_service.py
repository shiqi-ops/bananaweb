"""
New Mode service — 新模式（混合渲染：Dify 出内容 + 逐页 AI 渲染）后台任务编排。

流程:
    1. PENDING
    2. GENERATING  调用 Dify 工作流产出"内容"(每页标题/要点/视觉描述)
                   后端再用图片模型逐页渲染成高清图 → 组装成精美 PPTX
                   （未配置 Dify 时用本地占位内容，仍走同样的逐页渲染）
    3. DIAGNOSING  复用原 AI 诊断（DiagnosisTask + run_diagnosis_task）对成品逐页检查
    4. COMPLETED / FAILED  写入评分与摘要，前端可下载文件、查看报告
"""
import json
import logging
import os
import re
import shutil
from datetime import datetime

from models import db, NewModeTask, DiagnosisTask

logger = logging.getLogger(__name__)

# 无风格/场景时的兜底文案
DEFAULT_STYLE = '现代简洁商务风格，配色专业清晰，排版干净大方。'

_DIAGNOSIS_OPTIONS = ["layout", "color", "logic", "text"]


# ---------------------------------------------------------------------------
# 文件名 / 工具
# ---------------------------------------------------------------------------

def _sanitize_filename(name: str, max_len: int = 40) -> str:
    name = re.sub(r'[\\/:*?"<>|\r\n\t]', '', str(name or ''))
    name = name.strip()
    return (name[:max_len] or 'AI生成的演示文稿')


def _tmp_out_dir(app) -> str:
    d = os.path.join(app.config['UPLOAD_FOLDER'], 'new_mode')
    os.makedirs(d, exist_ok=True)
    return d


# ---------------------------------------------------------------------------
# 本地 PPTX 排版（python-pptx 文字模板风）
# ---------------------------------------------------------------------------

def _add_cover_slide(prs, title: str, style_hint: str):
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 顶部装饰条
    from pptx.enum.shapes import MSO_SHAPE
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.25))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(0xE8, 0x8A, 0x1A)
    bar.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.9), Inches(2.3), prs.slide_width - Inches(1.8), Inches(1.8))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title or '演示文稿'
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    p.font.name = 'Microsoft YaHei'

    if style_hint:
        sub = tf.add_paragraph()
        sub.text = style_hint
        sub.font.size = Pt(14)
        sub.font.color.rgb = RGBColor(0x8A, 0x94, 0xA6)
        sub.font.name = 'Microsoft YaHei'

    foot = slide.shapes.add_textbox(Inches(0.9), Inches(6.6), prs.slide_width - Inches(1.8), Inches(0.5))
    ft = foot.text_frame
    ft.paragraphs[0].text = '智绘视界 · AI 生成'
    ft.paragraphs[0].font.size = Pt(10)
    ft.paragraphs[0].font.color.rgb = RGBColor(0xB0, 0xB7, 0xC3)


def _add_content_slide(prs, page: dict, page_idx: int):
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    accent = RGBColor(0xE8, 0x8A, 0x1A)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(0.55), Inches(0.09), Inches(0.55))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    title = (page.get('title') or '').strip() or f'第 {page_idx} 页'
    tbox = slide.shapes.add_textbox(Inches(0.95), Inches(0.5), prs.slide_width - Inches(1.9), Inches(0.7))
    tf = tbox.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = title
    tf.paragraphs[0].font.size = Pt(26)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    tf.paragraphs[0].font.name = 'Microsoft YaHei'

    points = page.get('points') or page.get('bullets') or []
    if isinstance(points, str):
        points = [points]
    if not points:
        points = []
    body = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), prs.slide_width - Inches(2.0), Inches(5.2))
    btf = body.text_frame
    btf.word_wrap = True
    first = True
    for pt in points:
        text = str(pt).strip()
        if not text:
            continue
        # 兼容 '1. xxx' '- xxx' '• xxx'
        text = re.sub(r'^[\s]*([-•*]|\d+[.、)])\s*', '', text)
        if not text:
            continue
        para = btf.paragraphs[0] if first else btf.add_paragraph()
        first = False
        para.text = f'• {text}'
        para.font.size = Pt(16)
        para.font.color.rgb = RGBColor(0x3A, 0x44, 0x52)
        para.font.name = 'Microsoft YaHei'
        para.space_after = Pt(10)


def _pages_to_pptx(file_path: str, title: str, pages: list, style_hint: str = None) -> str:
    """把结构化 pages 排版成本地 PPTX 文件。"""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    _add_cover_slide(prs, title, style_hint or '')

    for i, page in enumerate(pages, start=1):
        _add_content_slide(prs, page, i)

    prs.save(file_path)
    return file_path


# ---------------------------------------------------------------------------
# 内容解析：Dify 返回 markdown / JSON → pages
# ---------------------------------------------------------------------------

def _parse_json_pages(obj, out: list):
    """从 JSON 结构中提取 pages: [{title, points}]。"""
    if isinstance(obj, list):
        for item in obj:
            _parse_json_pages(item, out)
        return
    if not isinstance(obj, dict):
        return
    for key in ('pages', 'slides', 'items', 'results'):
        if isinstance(obj.get(key), list):
            for item in obj[key]:
                _parse_json_pages(item, out)
    if 'title' in obj or 'heading' in obj or 'head' in obj:
        title = obj.get('title') or obj.get('heading') or obj.get('head') or ''
        body = obj.get('points') or obj.get('content') or obj.get('body') or obj.get('text') or []
        if isinstance(body, str):
            points = [l.strip() for l in re.split(r'[\n;；]', body) if l.strip()]
        elif isinstance(body, list):
            points = [str(x) for x in body if str(x).strip()]
        else:
            points = []
        if title or points:
            out.append({'title': str(title).strip(), 'points': points})


def _parse_markdown_pages(text: str) -> list:
    """markdown → pages。'#' 行视为页标题，其余要点为 bullets。"""
    lines = text.splitlines()
    pages: list = []
    cur_title = None
    cur_points: list = []
    heading_re = re.compile(r'^\s{0,3}#{1,6}\s+(.*)$')

    def flush():
        if cur_title or cur_points:
            pages.append({'title': cur_title or '', 'points': cur_points})
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        m = heading_re.match(line)
        if m:
            flush()
            cur_title = m.group(1).strip()
            cur_points = []
        else:
            # 也可能是 '第N页：xxx' 形式
            h2 = re.match(r'^第\s*\d+\s*[页Pp]:?\s*(.*)$', line)
            if h2:
                flush()
                cur_title = h2.group(1).strip()
                cur_points = []
            else:
                cur_points.append(line)
    flush()
    return pages


_META_HINTS = ('已深度思考', '深度思考', '用户输入', '输出到', '答案是', '所以', '直接输出',
               '我们输出', '推理', '因此', '综上', '结论是', '回答是', '下面是', '分析：', '分析:')


def _strip_meta_preamble(text: str) -> str:
    """剔除模型"深度思考/铺垫"式前缀行，避免污染第一页。

    只从开头剔除明显是 meta 说明的行；一旦遇到内容行（#、列表符、或普通短标题行）即停止。
    """
    lines = text.splitlines()
    start = 0
    found = False
    for i, raw in enumerate(lines):
        line = raw.strip()
        if not line:
            continue
        if line.startswith(('#', '-', '*', '•', '**')):
            start = i
            found = True
            break
        if any(h in line for h in _META_HINTS):
            continue
        start = i
        found = True
        break
    if not found:
        return text
    return '\n'.join(lines[start:])


def _parse_plain_text_pages(text: str) -> list:
    """纯文本 → pages：短标题行视为页标题，随后的句子视为该页要点。

    适配 Dify 输出"标题独占一行 + 普通句子"（无 #/## 标记）的情况。
    """
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    pages: list = []
    cur_title = None
    cur_points: list = []

    def flush():
        if cur_title or cur_points:
            pages.append({'title': cur_title or '', 'points': cur_points})

    for line in lines:
        # 标题行判定：较短、不以句末标点结尾、不是项目符号/编号/第N页
        is_title = (
            len(line) <= 24
            and not re.search(r'[。；！？，、：．.!?,:;]$', line)
            and not re.match(r'^[-•*#\d]', line)
            and not re.match(r'^第\s*\d+\s*[页Pp]', line)
        )
        if is_title:
            flush()
            cur_title = line
            cur_points = []
        else:
            cur_points.append(line)
    flush()
    return pages


def _parse_content_to_pages(text: str) -> list:
    text = (text or '').strip()
    if not text:
        return []
    stripped = text
    # 去掉可能的 ```json ... ``` 包装
    m = re.search(r'```(?:json)?\s*([\s\S]*?)```', text)
    if m:
        stripped = m.group(1).strip()
    # 剔除模型"深度思考/铺垫"式前缀
    stripped = _strip_meta_preamble(stripped)

    if stripped.startswith('{') or stripped.startswith('['):
        try:
            obj = json.loads(stripped)
            out: list = []
            _parse_json_pages(obj, out)
            if out:
                return out
        except (json.JSONDecodeError, TypeError):
            pass

    md_pages = _parse_markdown_pages(stripped)
    plain_pages = _parse_plain_text_pages(stripped)
    # 用页数更多的解析结果（纯文本能拆出标题行时优先）
    if len(plain_pages) > len(md_pages):
        return plain_pages
    if md_pages:
        return md_pages
    # 退化为单页
    return [{'title': '', 'points': [l for l in stripped.splitlines() if l.strip()]}]


# ---------------------------------------------------------------------------
# 占位实现（未配置 Dify 时使用，保证端到端可跑通）
# ---------------------------------------------------------------------------

_PLACEHOLDER_SECTIONS = [
    ('背景与意义', ['本演示围绕核心主题展开，先交代背景与要解决的问题。',
                  '说明该主题对目标受众的价值与必要性。']),
    ('核心思路', ['给出整体思路框架：分析现状 → 明确目标 → 制定方案 → 落地执行。',
                '每部分保持逻辑递进、重点突出。']),
    ('方案内容', ['阐述主要方案要点，并用简洁的条目分点说明。',
                '结合场景与风格要求组织语言。']),
    ('实施路径', ['拆解关键步骤与里程碑，方便后续推进。',
                '标注重点难点与风险应对。']),
    ('总结与展望', ['回顾核心结论与价值点。',
                  '提出下一步展望，收束整个演示。']),
]


def _placeholder_pages(requirement: str, page_count: int) -> (str, list):
    """根据需求生成占位页：首段尽力切分，不足用通用骨架补齐。"""
    lines = [l.strip() for l in (requirement or '').splitlines() if l.strip()]
    requirement_flat = ' '.join(lines)
    topic = lines[0] if lines else requirement_flat
    if len(topic) > 40:
        topic = topic[:40]

    pages: list = []
    # 需求里的每一段尽量作为一页要点
    para_idx = 1
    used = 0
    body_need = max(page_count - 1, 2)
    while used < body_need and lines:
        # 每次拿 1~2 行作为一页
        chunk = lines[:2]
        del lines[:2]
        points = []
        for c in chunk:
            if len(c) > 90:
                parts = re.split(r'(?<=[。；;！？])', c)
                points.extend([p for p in parts if p.strip()][:3])
            else:
                points.append(c)
        if not points:
            break
        pages.append({'title': f'需求要点 {para_idx}', 'points': points[:5]})
        para_idx += 1
        used += 1

    # 不足时用通用骨架补齐
    need = max(body_need - len(pages), 0)
    for i in range(need):
        sec_title, sec_points = _PLACEHOLDER_SECTIONS[i % len(_PLACEHOLDER_SECTIONS)]
        pages.append({'title': sec_title, 'points': list(sec_points)})
    pages = pages[:body_need]
    return topic, pages


def _placeholder_note() -> str:
    return ('占位模式：未检测到 DIFY_API_BASE / DIFY_API_KEY，已用本地模板生成一份可测试的 PPTX。'
            '配置 Dify 环境变量后即为真实「Dify 生成」。')


# ---------------------------------------------------------------------------
# 产出 PPTX（混合渲染：Dify 出内容 → 逐页 AI 渲染成图 → 组装 PPTX）
# ---------------------------------------------------------------------------

def _text_slide_image(title: str, points, style_hint: str, idx: int, out_dir: str) -> str:
    """兜底：用 PIL 本地画一页文字型幻灯片（图片模型失败时用）。"""
    from PIL import Image, ImageDraw, ImageFont

    w, h = 1920, 1080
    img = Image.new('RGB', (w, h), (250, 250, 252))
    draw = ImageDraw.Draw(img)
    font_path = os.path.join(os.path.dirname(__file__), '..', 'fonts', 'NotoSansSC-Regular.ttf')
    try:
        title_font = ImageFont.truetype(font_path, 72)
        body_font = ImageFont.truetype(font_path, 36)
    except Exception:
        title_font = ImageFont.load_default()
        body_font = ImageFont.load_default()

    draw.rectangle([(0, 0), (w, 26)], fill=(232, 138, 26))
    draw.text((120, 90), title or f'第 {idx} 页', fill=(31, 41, 55), font=title_font)
    y = 260
    for p in points or []:
        text = str(p).strip()
        if not text:
            continue
        text = re.sub(r'^[\s]*([-•*]|\d+[.、)])\s*', '', text)
        draw.text((150, y), f'• {text}', fill=(58, 68, 82), font=body_font)
        y += 62
    if style_hint:
        draw.text((150, y + 30), f'风格：{style_hint}', fill=(150, 150, 160), font=body_font)

    path = os.path.join(out_dir, f'page_{idx:02d}.png')
    img.convert('RGB').save(path)
    return path


def _render_one_slide(ai_service, title: str, points, style_hint: str, visual: str, idx: int, out_dir: str):
    """调用图片模型渲染单页高清图，成功返回图片路径，失败返回 None。"""
    # 过滤空标题与"第N页"这类序号标题，改用第一条要点当标题
    text_lines = []
    raw_title = str(title or '').strip()
    if raw_title and not re.match(r'^第\s*\d+\s*[页Pp]$', raw_title):
        text_lines.append(raw_title)
    for p in points or []:
        t = str(p).strip()
        if t:
            text_lines.append(t)

    content = "\n".join(text_lines) if text_lines else ""
    style = (style_hint or DEFAULT_STYLE).strip()
    visual_note = visual.strip() if visual else ""

    # 自绘一个干净的渲染提示词：明确"要渲染的文字"，并禁止模型渲染指令里的其他字样
    prompt = f"""你是一位顶级PPT设计师。请生成一张专业、美观、高清晰度的PPT页面。

【必须渲染的文字】（请把这些内容逐字、原样、不重不漏地作为本页的关键文字显示在页面上，第一行作为大标题，其余作为正文要点）：
{content}

【设计说明】
- 版式：大标题在顶部，下方用清晰的纵向要点列表（每条占一行），单栏布局。
- 整体风格：{style}
- 视觉/配图建议：{visual_note}
- 文字务必清晰锐利、易于阅读；每个要点不要写得太长。

【严格禁止】
- 禁止渲染本提示中的任何其它说明文字。
- 除非出现在【必须渲染的文字】里，否则页面上绝不能出现以下任何字眼：UI UX、4K、16:9、比例、分辨率、画质、画幅、清晰度、高清、PPT、演示文稿、设计师、第N页。
- 禁止把正文排成多列/多栏，禁止竖排文字，禁止拆分字词；正文用清晰中文，不要乱码、不要重复字。"""

    try:
        img = ai_service.generate_image(prompt, aspect_ratio='16:9', resolution='2K')
    except Exception as e:
        logger.warning(f"渲染第{idx}页图片失败，改用文字兜底: {e}")
        return None
    if img is None:
        logger.warning(f"渲染第{idx}页图片返回空，改用文字兜底")
        return None
    path = os.path.join(out_dir, f'page_{idx:02d}.png')
    try:
        img.convert('RGB').save(path)
        return path
    except Exception as e:
        logger.warning(f"保存第{idx}页图片失败: {e}")
        return None


def _ensure_content_pages(pages, target: int, cover_title: str = '') -> list:
    """把内容页数量凑到 target 页：不足则拆分要点或用通用骨架补足。"""
    pages = [p for p in (pages or []) if (p.get('title') or p.get('points') or p.get('bullets'))]
    if target <= 0:
        return pages
    if len(pages) >= target:
        return pages[:target]

    # 摊平所有要点，重新均分到 target 页
    all_points = []
    for p in pages:
        pts = p.get('points') or p.get('bullets') or []
        all_points.extend([str(x) for x in pts if str(x).strip()])

    if all_points:
        import math
        per = max(1, math.ceil(len(all_points) / target))
        out = []
        for i in range(target):
            chunk = all_points[i * per:(i + 1) * per]
            title = ''
            if i < len(pages) and (pages[i].get('title') or '').strip():
                title = str(pages[i]['title']).strip()
            else:
                title = f'第 {i + 1} 部分'
            out.append({'title': title, 'points': chunk or ['要点']})
        return out

    # 没有可用要点：用通用骨架补足
    sections = ['背景与意义', '核心思路', '方案内容', '实施路径', '总结与展望']
    out = []
    for i in range(target):
        s = sections[i % len(sections)]
        out.append({'title': s, 'points': [f'围绕「{cover_title or "主题"}」展开，说明{s}。']})
    return out


def _hybrid_render_pages(final_path: str, cover_title: str, pages, style_hint: str,
                         app, note: str, engine: str, page_count: int = None):
    """把结构化 pages 逐页渲染成图，再组装成精美 PPTX。返回 (engine, note)。

    page_count：目标总页数（含封面）；内容页不足时会自动拆分/补足到 page_count。
    """
    out_dir = os.path.join(app.config['UPLOAD_FOLDER'], 'new_mode', '_pages')
    os.makedirs(out_dir, exist_ok=True)

    from services.ai_service_manager import get_ai_service
    ai_service = get_ai_service()

    image_paths = []

    # 内容页；若第一页标题与封面相同则去掉，避免重复
    skip_first = bool(
        pages and pages[0].get('title') and cover_title
        and str(pages[0].get('title', '')).strip() == str(cover_title).strip()
    )
    content_pages = list(pages[1:] if skip_first else pages)

    # 决定是否单独加封面，以及内容页数量（目标：总页数 = page_count）
    use_cover = True
    if page_count and page_count > 1:
        if len(content_pages) >= page_count:
            # 内容已足够：不再额外加封面，直接取前 page_count 页
            use_cover = False
            content_pages = content_pages[:page_count]
        else:
            content_pages = _ensure_content_pages(content_pages, page_count - 1, cover_title or '')

    if use_cover:
        cover_path = _render_one_slide(ai_service, cover_title or '', [], style_hint, '', 1, out_dir)
        if cover_path is None:
            cover_path = _text_slide_image(cover_title or '演示文稿', [], style_hint, 1, out_dir)
        image_paths.append(cover_path)

    for idx, page in enumerate(content_pages, start=1):
        title = str(page.get('title') or '').strip() or f'第 {idx} 页'
        points = page.get('points') or page.get('bullets') or []
        visual = str(page.get('visual') or page.get('image_desc') or page.get('layout') or '').strip()
        p = _render_one_slide(ai_service, title, points, style_hint, visual, idx, out_dir)
        if p is None:
            p = _text_slide_image(title, points, style_hint, idx, out_dir)
        image_paths.append(p)

    if not image_paths:
        raise RuntimeError('所有页面的图片渲染均失败，请检查图片模型配置')

    from services.export_service import ExportService
    ExportService().create_pptx_from_images(image_paths, output_file=final_path, aspect_ratio='16:9')
    return engine, note


def _produce_pptx(task, final_path: str, app):
    """生成/获取成品 PPTX 到 final_path，返回 (engine, note)。

    新模式 = 「混合渲染」：Dify 出内容 → 图片模型逐页渲染成高清图 → 组装 PPTX。
    """
    title = (task.title or '').strip() or _sanitize_filename((task.requirement or '演示文稿')[:40], 40)
    style_hint = (task.style_description or DEFAULT_STYLE).strip()
    inputs = {
        'requirement': task.requirement or '',
        'title': title,
        'page_count': task.page_count or 10,
        'usage_scenario': task.usage_scenario or '',
        'style_description': style_hint,
    }

    from services.dify_service import is_dify_configured, run_dify_workflow
    tmp_path = None
    if not is_dify_configured():
        topic, pages = _placeholder_pages(task.requirement or '', task.page_count or 10)
        return _hybrid_render_pages(final_path, topic or title, pages, style_hint, app,
                                    note=_placeholder_note(), engine='placeholder',
                                    page_count=task.page_count or 10)

    try:
        result = run_dify_workflow(inputs)
    except Exception as e:
        logger.error(f"新模式 Dify 调用失败，降级为占位生成: {e}")
        topic, pages = _placeholder_pages(task.requirement or '', task.page_count or 10)
        return _hybrid_render_pages(final_path, topic or title, pages, style_hint, app,
                                    note=f'Dify 调用失败已降级为占位生成：{e}', engine='placeholder',
                                    page_count=task.page_count or 10)

    try:
        text = result.get('text') or ''
        pages = _parse_content_to_pages(text) if text else []
        if pages:
            # 优先：内容 → 逐页 AI 渲染（更美观）
            return _hybrid_render_pages(final_path, title, pages, style_hint, app,
                                        note='由 Dify 生成内容，逐页 AI 渲染为精美 PPTX', engine='dify',
                                        page_count=task.page_count or 10)
        # 兜底：Dify 直接返回了 .pptx 文件（无可用内容时）
        if result.get('source') == 'file':
            tmp_path = result.get('file_path')
            if tmp_path and os.path.exists(tmp_path):
                shutil.move(tmp_path, final_path)
                tmp_path = None
                return 'dify', '由 Dify 工作流直接生成成品 PPTX'
        raise RuntimeError('Dify 返回的内容无法解析出页面结构')
    except Exception as e:
        logger.error(f"处理 Dify 产出失败: {e}")
        raise RuntimeError(f'处理 Dify 产出失败: {e}') from e
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.unlink(tmp_path)
            except Exception:
                pass


# ---------------------------------------------------------------------------
# 后台任务入口
# ---------------------------------------------------------------------------

def run_new_mode_task(task_id: str, app):
    """
    后台执行新模式生成任务（由 TaskManager.submit_task 调用）。

    参数与 TaskManager 约定一致: func(task_id, *args)。
    """
    if app is None:
        raise ValueError("Flask app instance must be provided")

    with app.app_context():
        task = None
        try:
            task = NewModeTask.get_by_id(task_id)
            if task is None:
                logger.error(f"新模式任务 {task_id} 不存在")
                return

            # ---- 1. 生成 ----
            task.status = 'GENERATING'
            task.note = '正在通过 Dify 工作流生成 PPT...' if task.engine == 'dify' else '正在生成 PPT...'
            task.engine = 'dify'  # 生成引擎统一标识为 dify（未配置时实为占位）
            db.session.commit()
            logger.info(f"新模式 {task_id}: 开始生成")

            final_dir = _tmp_out_dir(app)
            final_path = os.path.join(final_dir, f"{task.id}.pptx")
            engine, note = _produce_pptx(task, final_path, app)
            if not os.path.exists(final_path):
                raise RuntimeError('PPTX 生成失败：未产出文件')

            task.engine = engine
            task.note = note
            task.file_path = final_path
            task.file_name = f"{_sanitize_filename(task.title or 'AI生成演示文稿', 40)}.pptx"
            db.session.commit()
            logger.info(f"新模式 {task_id}: PPTX 已生成 ({engine})")

            # ---- 2. 复用原 AI 诊断 ----
            task.status = 'DIAGNOSING'
            task.note = '正在使用 AI 对生成的 PPT 逐页诊断...'
            db.session.commit()

            diagnosis = DiagnosisTask(
                user_id=task.user_id,
                file_path=final_path,
                file_type='pptx',
                diagnosis_options=json.dumps(_DIAGNOSIS_OPTIONS, ensure_ascii=False),
            )
            if not DiagnosisTask.create_task(diagnosis):
                raise RuntimeError('创建诊断任务失败（数据库写入错误）')
            task.diagnosis_task_id = diagnosis.id
            db.session.commit()

            from services.diagnosis_service import run_diagnosis_task
            run_diagnosis_task(diagnosis.id, app)

            diag = DiagnosisTask.get_by_id(diagnosis.id)
            note_suffix = ''
            if diag and diag.status == 'COMPLETED':
                try:
                    result = json.loads(diag.result) if diag.result else {}
                    task.score = result.get('score')
                    task.summary = result.get('summary')
                    task.result = diag.result
                    note_suffix = f'（AI 诊断评分 {result.get("score")}/100）'
                except (json.JSONDecodeError, TypeError):
                    pass
            elif diag and diag.status == 'FAILED':
                note_suffix = f'（AI 诊断未完成: {diag.error_message or "未知错误"}，文件仍可下载）'

            task.status = 'COMPLETED'
            task.note = f"{note}{note_suffix}"
            task.completed_at = datetime.utcnow()
            db.session.commit()
            logger.info(f"新模式 {task_id} 完成: {task.note}")

        except Exception as e:
            logger.error(f"新模式任务 {task_id} 失败: {e}", exc_info=True)
            try:
                db.session.rollback()
                task = NewModeTask.get_by_id(task_id) if task_id else None
                if task:
                    task.status = 'FAILED'
                    task.error_message = str(e)
                    task.completed_at = datetime.utcnow()
                    db.session.commit()
            except Exception as db_err:
                logger.error(f"更新新模式失败状态出错: {db_err}")
                db.session.rollback()
